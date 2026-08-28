# -*- coding: utf-8 -*-
"""Nhập BoQ mời thầu từ file Excel khách gửi.

Không có bước này thì module chỉ dùng được trên lý thuyết: gói thầu thật
đầu tiên đã có 163 dòng, không ai gõ tay.

VÌ SAO ĐỌC "MỀM" CHỨ KHÔNG ÉP KHUÔN CỐ ĐỊNH:

File mời thầu do khách soạn, mỗi khách một kiểu — tiêu đề nằm ở dòng 3
hay dòng 6, cột "Khối lượng" có thể tên "KL" hoặc "Số lượng", trên đầu
thường có mấy dòng thư ngỏ. Ép một khuôn cứng thì gần như file nào cũng
phải sửa tay trước khi nhập, và người dùng sẽ quay về gõ tay cho nhanh.

Nên: tự dò dòng tiêu đề bằng từ khoá, cho người dùng xem trước và sửa
số dòng tiêu đề nếu dò sai. Dò sai còn chữa được; đọc nhầm im lặng thì
không.

DÒNG TIÊU ĐỀ NHÓM (A, B, C…) VẪN NHẬP: chúng không có khối lượng nhưng
là bộ khung của file, và hồ sơ nộp lại phải giữ đúng bộ khung đó.
"""
import base64
import io
import logging

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# từ khoá nhận diện cột — so khớp không dấu, không phân biệt hoa thường
COL_HINTS = {
    'code': ['stt', 'ma', 'ky hieu'],
    'name': ['noi dung', 'ten cong viec', 'hang muc', 'mo ta', 'cong viec'],
    'uom': ['don vi', 'dvt', 'dv tinh'],
    'qty': ['khoi luong', 'kl', 'so luong'],
}


def _norm(text):
    """Bỏ dấu tiếng Việt và chuẩn hoá để so khớp tiêu đề cột."""
    if not text:
        return ''
    s = str(text).strip().lower()
    pairs = (
        ('àáạảãâầấậẩẫăằắặẳẵ', 'a'), ('èéẹẻẽêềếệểễ', 'e'),
        ('ìíịỉĩ', 'i'), ('òóọỏõôồốộổỗơờớợởỡ', 'o'),
        ('ùúụủũưừứựửữ', 'u'), ('ỳýỵỷỹ', 'y'), ('đ', 'd'),
    )
    for chars, repl in pairs:
        for c in chars:
            s = s.replace(c, repl)
    return ' '.join(s.split())


class RpBidImport(models.TransientModel):
    _name = 'rp.bid.import'
    _description = 'Nhập BoQ mời thầu từ Excel'

    bid_id = fields.Many2one(
        'rp.bid', string='Hồ sơ dự thầu', required=True,
        ondelete='cascade')
    file_data = fields.Binary(string='File Excel', required=True)
    file_name = fields.Char(string='Tên file')
    sheet_name = fields.Char(
        string='Sheet', help='Bỏ trống = lấy sheet đầu tiên.')
    header_row = fields.Integer(
        string='Dòng tiêu đề', default=0,
        help='Để 0 để hệ thống tự dò. Dò sai thì điền số dòng tiêu đề vào '
             'đây rồi xem trước lại.')
    replace_existing = fields.Boolean(
        string='Xoá dòng chào giá cũ trước khi nhập', default=True)

    preview = fields.Text(string='Xem trước', readonly=True)
    state = fields.Selection(
        [('upload', 'Chọn file'), ('preview', 'Xem trước')],
        default='upload')

    # ==================================================================
    def _load_sheet(self):
        try:
            import openpyxl
        except ImportError:
            raise UserError(_('Máy chủ thiếu thư viện openpyxl.'))
        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.file_data)),
                data_only=True, read_only=True)
        except Exception as exc:
            raise UserError(_(
                'Không đọc được file: %s\n\nFile phải là .xlsx / .xlsm. '
                'File .xls đời cũ thì mở bằng Excel và lưu lại dạng .xlsx.',
                exc))
        if self.sheet_name and self.sheet_name in wb.sheetnames:
            return wb, wb[self.sheet_name]
        return wb, wb.worksheets[0]

    def _detect_header(self, rows):
        """Tìm dòng tiêu đề: dòng khớp nhiều từ khoá cột nhất."""
        best, best_score = 0, 0
        for idx, row in enumerate(rows[:30]):
            cells = [_norm(c) for c in row]
            score = 0
            for hints in COL_HINTS.values():
                if any(any(h == c or h in c for h in hints) for c in cells):
                    score += 1
            if score > best_score:
                best, best_score = idx, score
        if best_score < 3:
            raise UserError(_(
                'Không dò được dòng tiêu đề. Mở file xem tiêu đề nằm ở dòng '
                'thứ mấy rồi điền vào ô "Dòng tiêu đề" và bấm Xem trước lại.'))
        return best

    def _map_columns(self, header):
        cols = {}
        for idx, cell in enumerate(header):
            norm = _norm(cell)
            if not norm:
                continue
            for key, hints in COL_HINTS.items():
                if key in cols:
                    continue
                if any(h == norm or norm.startswith(h) for h in hints):
                    cols[key] = idx
        missing = [k for k in ('name', 'uom', 'qty') if k not in cols]
        if 'name' in missing:
            raise UserError(_(
                'Không tìm thấy cột nội dung công việc trong dòng tiêu đề.'))
        return cols

    def _parse(self):
        self.ensure_one()
        wb, ws = self._load_sheet()
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        if not rows:
            raise UserError(_('Sheet rỗng.'))

        hdr = (self.header_row - 1) if self.header_row else self._detect_header(rows)
        if hdr < 0 or hdr >= len(rows):
            raise UserError(_('Dòng tiêu đề nằm ngoài phạm vi sheet.'))
        cols = self._map_columns(rows[hdr])

        def cell(row, key):
            i = cols.get(key)
            return row[i] if i is not None and i < len(row) else None

        items, seq = [], 0
        for row in rows[hdr + 1:]:
            name = cell(row, 'name')
            if not name or not str(name).strip():
                continue
            name = ' '.join(str(name).split())
            qty = cell(row, 'qty')
            qty = float(qty) if isinstance(qty, (int, float)) else 0.0
            uom = cell(row, 'uom')
            uom = str(uom).strip() if uom else ''
            code = cell(row, 'code')
            code = str(code).strip() if code not in (None, '') else ''
            # DỪNG HẲN ở dòng tổng cộng, không chỉ bỏ qua nó.
            #
            # Phiếu mời thầu thật gần như luôn có phần văn bản phía dưới
            # bảng — "Trách nhiệm Bên A", "Trách nhiệm Bên B", điều kiện
            # thanh toán — và các đoạn đó cũng đánh số 1, 2, 3 nên trông
            # y hệt dòng BoQ. Chỉ bỏ qua dòng tổng thì 14 đoạn văn của
            # file đầu tiên đã lọt vào thành dòng chào giá.
            #
            # Bảng khối lượng kết thúc tại dòng tổng — sau đó không còn
            # gì để đọc nữa.
            norm_name = _norm(name)
            if norm_name.startswith(('tong cong', 'cong truoc thue',
                                     'tong gia tri', 'tong')):
                break
            if norm_name.startswith('thue'):
                continue
            # Dòng tiêu đề nhóm: có mã/nội dung nhưng không có khối lượng.
            is_section = not qty and not uom
            seq += 1
            items.append({
                'sequence': seq * 10, 'code': code, 'name': name[:400],
                'uom_text': uom, 'quantity_client': qty,
                'is_section': is_section,
            })
        if not items:
            raise UserError(_(
                'Không đọc được dòng nào bên dưới tiêu đề. Kiểm lại số dòng '
                'tiêu đề.'))
        return items

    # ==================================================================
    def action_preview(self):
        self.ensure_one()
        items = self._parse()
        head = items[:15]
        lines = ['%-6s %-52s %-10s %14s' % ('MÃ', 'NỘI DUNG', 'ĐVT', 'KL')]
        lines.append('-' * 86)
        for it in head:
            lines.append('%-6s %-52s %-10s %14s' % (
                it['code'][:6], it['name'][:52], it['uom_text'][:10],
                '' if it['is_section'] else '{:,.3f}'.format(it['quantity_client'])))
        n_sec = len([i for i in items if i['is_section']])
        lines.append('')
        lines.append('Tổng %d dòng (%d dòng tiêu đề nhóm, %d dòng có khối lượng)'
                     % (len(items), n_sec, len(items) - n_sec))
        if len(items) > 15:
            lines.append('… còn %d dòng nữa' % (len(items) - 15))
        self.preview = '\n'.join(lines)
        self.state = 'preview'
        return {'type': 'ir.actions.act_window', 'res_model': self._name,
                'res_id': self.id, 'view_mode': 'form', 'target': 'new'}

    def action_import(self):
        self.ensure_one()
        items = self._parse()
        if self.replace_existing:
            self.bid_id.item_ids.unlink()
        self.env['rp.bid.item'].create(
            [dict(it, bid_id=self.bid_id.id) for it in items])
        _logger.info('rp_bid: nhập %s dòng chào giá vào hồ sơ %s',
                     len(items), self.bid_id.code or self.bid_id.id)
        return self.bid_id.action_open_items()
