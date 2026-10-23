# -*- coding: utf-8 -*-
"""Đọc tệp nguồn — phần duy nhất biết về openpyxl và csv.

Gom mọi thứ phụ thuộc thư viện đọc tệp vào một chỗ, để phần nghiệp vụ
chỉ làm việc với *địa chỉ → giá trị*. Thêm định dạng mới sau này là thêm
một hàm ở đây, không đụng phần ghi dữ liệu.

Vùng trộn — chỗ sai im lặng nhất khi khai địa chỉ ô
---------------------------------------------------
Trong tờ báo cáo thật, nhãn ``Plant Irradiance`` trộn từ ``L19`` tới
``O19`` nên giá trị nằm ở ``P19``; nhưng nhãn ``Energy loss for
Curtailment Outage`` ngay bên dưới chỉ trộn tới ``N30`` nên giá trị nằm
ở ``O30``. Cùng một khối, cùng cột nhãn, hai cột giá trị khác nhau.

Người khai mẫu sẽ trỏ nhầm, và openpyxl trả về ``None`` cho mọi ô nằm
trong vùng trộn mà không phải góc trên trái — không lỗi, không cảnh
báo, trường chỉ rỗng. Nên ``doc_o`` trả về kèm *gợi ý* ô thật khi gặp
trường hợp đó.
"""
import csv
import io
import re
from datetime import date, datetime, time

NGAY_DANG = ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d%m%Y')

RE_O = re.compile(r'^([A-Za-z]{1,3})([0-9]{1,7})$')
RE_VUNG = re.compile(r'^([A-Za-z]{1,3}[0-9]{1,7}):([A-Za-z]{1,3}[0-9]{1,7})$')


class LoiNguon(Exception):
    """Tệp hoặc địa chỉ không đọc được — kèm câu giải thích cho người dùng."""


# ======================================================================
# Địa chỉ
# ======================================================================
def tach_o(dia_chi):
    """``'B20'`` → ``('B', 20)``. Ném lỗi nếu không phải địa chỉ ô."""
    m = RE_O.match((dia_chi or '').strip())
    if not m:
        raise LoiNguon(
            'Địa chỉ ô "%s" không hợp lệ — phải dạng B20, AC7.' % dia_chi)
    return m.group(1).upper(), int(m.group(2))


def tach_vung(dia_chi):
    """``'D40:P40'`` → danh sách địa chỉ ô theo thứ tự đọc.

    Chỉ nhận vùng **một dòng** hoặc **một cột**: dòng giờ trong tờ báo
    cáo là một dải ngang, và một vùng hai chiều thì không có thứ tự duy
    nhất để ghép với dải giá trị tương ứng.
    """
    from openpyxl.utils import column_index_from_string, get_column_letter
    m = RE_VUNG.match((dia_chi or '').strip())
    if not m:
        raise LoiNguon(
            'Vùng "%s" không hợp lệ — phải dạng D40:P40.' % dia_chi)
    c1, r1 = tach_o(m.group(1))
    c2, r2 = tach_o(m.group(2))
    i1, i2 = column_index_from_string(c1), column_index_from_string(c2)
    if r1 != r2 and i1 != i2:
        raise LoiNguon(
            'Vùng "%s" trải cả hai chiều. Chỉ nhận một dòng hoặc một cột, '
            'vì vùng hai chiều không có thứ tự duy nhất để ghép với dải '
            'giá trị.' % dia_chi)
    if r1 == r2:
        return ['%s%d' % (get_column_letter(i), r1)
                for i in range(min(i1, i2), max(i1, i2) + 1)]
    return ['%s%d' % (c1, r) for r in range(min(r1, r2), max(r1, r2) + 1)]


# ======================================================================
# Excel — tờ biểu
# ======================================================================
def mo_workbook(noi_dung):
    try:
        import openpyxl
    except ImportError:                                  # pragma: no cover
        raise LoiNguon('Thiếu thư viện openpyxl trên máy chủ.')
    try:
        return openpyxl.load_workbook(
            io.BytesIO(noi_dung), data_only=True, read_only=False)
    except Exception as e:
        raise LoiNguon(
            'Không mở được tệp Excel: %s. Tệp .xls cũ phải lưu lại thành '
            '.xlsx.' % (str(e)[:160],))


def ban_do_tron(ws):
    """``{địa chỉ ô bị phủ: địa chỉ góc trên trái}`` cho mọi vùng trộn."""
    bd = {}
    for v in ws.merged_cells.ranges:
        goc = v.start_cell.coordinate
        for dong in v.cells:
            from openpyxl.utils import get_column_letter
            o = '%s%d' % (get_column_letter(dong[1]), dong[0])
            if o != goc:
                bd[o] = goc
    return bd


def doc_o(ws, dia_chi, tron=None):
    """Giá trị tại một ô, kèm gợi ý khi ô nằm trong vùng trộn.

    Trả về ``(giá trị, ghi chú)``. Ghi chú khác rỗng nghĩa là địa chỉ
    khai có vấn đề đáng nói — giá trị vẫn trả về như openpyxl thấy
    (thường là ``None``) để người khai thấy đúng hậu quả của địa chỉ
    mình gõ, chứ không bị trình nạp âm thầm "sửa hộ".
    """
    cot, dong = tach_o(dia_chi)
    o = '%s%d' % (cot, dong)
    gt = ws[o].value
    if gt is None and tron and o in tron:
        return None, ('ô %s nằm trong vùng trộn, giá trị thật ở ô %s'
                      % (o, tron[o]))
    return gt, ''


# ======================================================================
# Bảng — CSV hoặc Excel một dòng tiêu đề
# ======================================================================
def doc_bang(noi_dung, ten_tep):
    """Đọc tệp dạng bảng thành ``(tiêu đề, danh sách dòng)``.

    Dòng là ``list`` giá trị thô theo đúng thứ tự tiêu đề.
    """
    ten = (ten_tep or '').lower()
    if ten.endswith('.csv') or ten.endswith('.txt'):
        return _doc_csv(noi_dung)
    wb = mo_workbook(noi_dung)
    ws = wb[wb.sheetnames[0]]
    dong = list(ws.iter_rows(values_only=True))
    if not dong:
        raise LoiNguon('Tệp không có dòng nào.')
    tieu_de = [str(c).strip() if c is not None else '' for c in dong[0]]
    return tieu_de, [list(d) for d in dong[1:]]


def _doc_csv(noi_dung):
    van = None
    for bo_ma in ('utf-8-sig', 'utf-8', 'cp1258', 'latin-1'):
        try:
            van = noi_dung.decode(bo_ma)
            break
        except UnicodeDecodeError:
            continue
    if van is None:                                      # pragma: no cover
        raise LoiNguon('Không đoán được bộ mã của tệp CSV.')
    mau = van[:4096]
    try:
        dau = csv.Sniffer().sniff(mau, delimiters=',;\t|').delimiter
    except csv.Error:
        # Sniffer thất bại với tệp một cột hoặc tiêu đề lạ. Phẩy là mặc
        # định đúng trong phần lớn trường hợp, và dòng nào tách sai sẽ
        # hiện ra ở nhật ký chứ không làm hỏng cả lượt nạp.
        dau = ','
    doc = csv.reader(io.StringIO(van), delimiter=dau)
    dong = list(doc)
    if not dong:
        raise LoiNguon('Tệp CSV không có dòng nào.')
    return [c.strip() for c in dong[0]], dong[1:]


# ======================================================================
# Chuyển kiểu
# ======================================================================
def ra_so(gt):
    """Giá trị ô → float, hoặc ``None`` nếu không phải số.

    Nhận cả chuỗi kiểu ``"1.234,56"`` (dấu nghìn kiểu Việt) và
    ``"1,234.56"`` (kiểu Anh) — tệp xuất từ hệ của hãng dùng cả hai, và
    đoán sai thì ra con số lệch một nghìn lần.
    """
    if gt is None or isinstance(gt, bool):
        return None
    if isinstance(gt, (int, float)):
        return float(gt)
    s = str(gt).strip().replace('\xa0', '').replace(' ', '')
    if not s:
        return None
    s = s.rstrip('%')
    if ',' in s and '.' in s:
        # Dấu nào ở sau cùng thì dấu đó là dấu thập phân.
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        phan = s.split(',')
        # "1,234" là một nghìn hai trăm; "6,36" là sáu phẩy ba sáu. Phân
        # biệt bằng số chữ số sau dấu — đúng ba thì là dấu nghìn.
        s = s.replace(',', '') if len(phan[-1]) == 3 else s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return None


def ra_ngay(gt):
    """Giá trị ô → ``date``, hoặc ``None``."""
    if isinstance(gt, datetime):
        return gt.date()
    if isinstance(gt, date):
        return gt
    if gt is None:
        return None
    s = str(gt).strip()
    if not s:
        return None
    for d in NGAY_DANG:
        try:
            return datetime.strptime(s, d).date()
        except ValueError:
            continue
    return None


def ra_gio(gt):
    """Giá trị ô → số giờ trong ngày (float), hoặc ``None``.

    Ô giờ trong Excel về tới đây dưới ba dạng khác nhau tuỳ cách ô được
    định dạng: ``time``, ``datetime``, hay phân số của một ngày.
    """
    if isinstance(gt, time):
        return gt.hour + gt.minute / 60.0 + gt.second / 3600.0
    if isinstance(gt, datetime):
        return gt.hour + gt.minute / 60.0 + gt.second / 3600.0
    so = ra_so(gt)
    if so is None:
        return None
    # Phân số của một ngày (0,5 = 12h) so với số giờ thẳng (13 = 13h).
    return so * 24.0 if 0 <= so <= 1 else so


def ngay_tu_ten_sheet(ten):
    """Rút ngày từ tên sheet kiểu ``21102026`` hoặc ``2026-10-21``.

    Trả về ``None`` khi tên sheet không mang ngày — đó là chuyện bình
    thường (sheet tên "Sheet1"), không phải lỗi.
    """
    s = (ten or '').strip()
    for d in ('%d%m%Y', '%Y%m%d', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, d).date()
        except ValueError:
            continue
    m = re.search(r'(\d{8})', s)
    if m:
        for d in ('%d%m%Y', '%Y%m%d'):
            try:
                return datetime.strptime(m.group(1), d).date()
            except ValueError:
                continue
    return None
